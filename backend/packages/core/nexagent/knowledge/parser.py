"""Document parser — extract plain text from PDF, DOCX, PPTX, XLSX, TXT, MD, HTML.

All parsing is async-friendly: heavy work is delegated to a thread pool so it
doesn't block the event loop.
"""

from __future__ import annotations

import asyncio
import hashlib
import logging
import posixpath
import re
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

logger = logging.getLogger(__name__)


@dataclass
class Table:
    markdown: str
    metadata: dict = field(default_factory=dict)


@dataclass
class Image:
    path: str
    alt: str = ""
    metadata: dict = field(default_factory=dict)


@dataclass
class ParseConfig:
    use_ocr: bool = False
    prefer_docling: bool = True


@dataclass
class ParsedDocument:
    content: str
    tables: list[Table] = field(default_factory=list)
    images: list[Image] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


class DocumentParser:
    """Unified parser facade returning Markdown plus metadata."""

    def parse(self, file_path: str | Path, config: ParseConfig | None = None) -> ParsedDocument:
        path = Path(file_path)
        config = config or ParseConfig()
        suffix = path.suffix.lower()
        parser_chain: list[dict] = []

        if config.prefer_docling and suffix in {".pdf", ".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".xls"}:
            try:
                parsed = self._parse_with_docling(path)
                parsed.metadata["parser_chain"] = [{"engine": "docling", "status": "ok"}]
                return parsed
            except Exception as exc:
                parser_chain.append({"engine": "docling", "status": "failed", "reason": str(exc)})
                logger.debug("Docling parser failed for %s; falling back: %s", path, exc)

        metadata_extra: dict = {}
        match suffix:
            case ".pdf":
                try:
                    content, metadata_extra = _parse_pdf_with_pymupdf(path)
                    parser_chain.append({"engine": "pymupdf", "status": "ok"})
                except Exception as exc:
                    parser_chain.append({"engine": "pymupdf", "status": "failed", "reason": str(exc)})
                    content = _parse_pdf(str(path))
                    parser_chain.append({"engine": "pypdf", "status": "ok"})
            case ".docx" | ".doc":
                content = _parse_docx(str(path))
                parser_chain.append({"engine": "docx", "status": "ok"})
            case ".pptx" | ".ppt":
                content = _parse_pptx(str(path))
                parser_chain.append({"engine": "pptx", "status": "ok"})
            case ".xlsx":
                content, metadata_extra = _parse_xlsx(path)
                parser_chain.append({"engine": "xlsx", "status": "ok"})
            case ".html" | ".htm":
                content = _parse_html(str(path))
                parser_chain.append({"engine": "html", "status": "ok"})
            case ".png" | ".jpg" | ".jpeg" | ".bmp" | ".tiff" | ".tif":
                content = _parse_image(str(path), use_ocr=config.use_ocr)
                parser_chain.append({"engine": "ocr", "status": "ok"})
            case ".md" | ".markdown" | ".txt" | ".rst" | ".csv" | ".json" | ".yaml" | ".yml":
                content = _parse_text(str(path))
                parser_chain.append({"engine": "text", "status": "ok"})
            case _:
                content = _parse_text(str(path))
                parser_chain.append({"engine": "text", "status": "ok"})

        metadata = _base_metadata(path, content, "fallback")
        if parser_chain:
            metadata["parser_chain"] = parser_chain
            ok_engines = [item["engine"] for item in parser_chain if item.get("status") == "ok"]
            if ok_engines:
                metadata["parser"] = ok_engines[-1]
        metadata.update(metadata_extra)
        return ParsedDocument(
            content=content,
            metadata=metadata,
        )

    def _parse_with_docling(self, path: Path) -> ParsedDocument:
        try:
            from docling.document_converter import DocumentConverter  # type: ignore
        except ImportError as exc:
            raise RuntimeError("docling is not installed") from exc

        result = DocumentConverter().convert(path)
        document = result.document
        markdown = document.export_to_markdown()
        return ParsedDocument(
            content=markdown,
            metadata={
                **_base_metadata(path, markdown, "docling"),
                "parser": "docling",
                "status": getattr(getattr(result, "status", None), "name", str(getattr(result, "status", ""))),
            },
        )


async def parse_document(file_path: str) -> str:
    """Extract plain text from *file_path* and return it as a string.

    Dispatches by file extension.  Falls back to UTF-8 read for unknown types.
    Runs blocking I/O in a thread-pool executor.
    """
    loop = asyncio.get_event_loop()
    parsed = await loop.run_in_executor(None, DocumentParser().parse, file_path, ParseConfig())
    return parsed.content


async def parse_document_structured(file_path: str, config: ParseConfig | None = None) -> ParsedDocument:
    """Parse a document and preserve parser metadata for downstream evidence."""
    loop = asyncio.get_event_loop()
    return await loop.run_in_executor(None, DocumentParser().parse, file_path, config or ParseConfig())


# ── per-format parsers (synchronous, run in thread pool) ─────────────────────

def _parse_pdf(file_path: str) -> str:
    """Extract text from a PDF using pypdf."""
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        raise ImportError("pypdf is required to parse PDF files. Run: pip install pypdf")

    reader = PdfReader(file_path)
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            parts.append(text.strip())
    return "\n\n".join(parts)


def _parse_pdf_with_pymupdf(path: Path) -> tuple[str, dict]:
    """Extract a PDF as Markdown-ish text with page markers and section hints."""
    try:
        import fitz  # type: ignore
    except ImportError as exc:
        raise ImportError("PyMuPDF is not installed") from exc

    doc = fitz.open(str(path))
    page_parts: list[str] = []
    title = ""
    headings: list[str] = []
    for page_index, page in enumerate(doc, 1):
        text = page.get_text("text") or ""
        lines = [_clean_pdf_line(line) for line in text.splitlines()]
        lines = [line for line in lines if line]
        if not lines:
            continue
        if not title:
            title = _guess_pdf_title(lines)
        rendered: list[str] = [f"<!-- page:{page_index} -->"]
        for line in lines:
            normalized = _promote_pdf_heading(line)
            if normalized.startswith("## "):
                headings.append(normalized[3:].strip())
            rendered.append(normalized)
        page_parts.append("\n".join(rendered))
    doc.close()
    content = "\n\n".join(page_parts)
    return content, {
        "parser": "pymupdf",
        "page_count": len(page_parts),
        "title": title,
        "headings": headings[:100],
        "degraded": False,
    }


def _clean_pdf_line(line: str) -> str:
    line = re.sub(r"\s+", " ", line or "").strip()
    return line.replace("\x00", "")


def _guess_pdf_title(lines: list[str]) -> str:
    candidates = [line for line in lines[:12] if 8 <= len(line) <= 180]
    if not candidates:
        return ""
    return max(candidates, key=len)


def _promote_pdf_heading(line: str) -> str:
    heading_re = re.compile(
        r"^(?:\d+(?:\.\d+)*\s+)?"
        r"(Abstract|Introduction|Related Work|Background|Method|Methodology|Approach|"
        r"Experiments?|Evaluation|Results?|Discussion|Conclusion|References|Bibliography|"
        r"摘要|引言|相关工作|方法|实验|结果|讨论|结论|参考文献)\b",
        re.IGNORECASE,
    )
    if line.startswith("#"):
        return line
    if heading_re.match(line) and len(line) <= 120:
        return f"## {line}"
    return line


def _parse_docx(file_path: str) -> str:
    """Extract text from a .docx file using python-docx."""
    try:
        from docx import Document  # type: ignore
    except ImportError:
        raise ImportError("python-docx is required to parse DOCX files. Run: pip install python-docx")

    doc = Document(file_path)
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    # Also extract table cells
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts)


def _parse_pptx(file_path: str) -> str:
    """Extract text from a .pptx file using python-pptx."""
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        raise ImportError("python-pptx is required to parse PPTX files. Run: pip install python-pptx")

    prs = Presentation(file_path)
    parts: list[str] = []
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_texts.append(shape.text.strip())
        if slide_texts:
            parts.append(f"[Slide {slide_num}]\n" + "\n".join(slide_texts))
    return "\n\n".join(parts)


def _parse_xlsx(path: Path) -> tuple[str, dict]:
    """Extract workbook rows from .xlsx files as Markdown tables."""
    try:
        with ZipFile(path) as archive:
            shared_strings = _xlsx_shared_strings(archive)
            sheets = _xlsx_sheet_refs(archive)
            parts: list[str] = []
            total_rows = 0
            for sheet_index, (sheet_name, sheet_path) in enumerate(sheets, 1):
                try:
                    xml = archive.read(sheet_path)
                except KeyError:
                    continue
                rows = _xlsx_rows(xml, shared_strings)
                total_rows += len(rows)
                if rows:
                    parts.append(f"## Sheet: {sheet_name}\n\n{_xlsx_rows_to_markdown(rows)}")
                else:
                    parts.append(f"## Sheet: {sheet_name}\n\n(empty)")
            content = "\n\n".join(parts)
            return content, {
                "parser": "xlsx",
                "sheet_count": len(sheets),
                "row_count": total_rows,
                "degraded": False,
            }
    except BadZipFile as exc:
        raise ValueError(f"{path.name} is not a valid .xlsx file.") from exc
    except ET.ParseError as exc:
        raise ValueError(f"{path.name} contains invalid spreadsheet XML.") from exc


def _xlsx_shared_strings(archive: ZipFile) -> list[str]:
    try:
        xml = archive.read("xl/sharedStrings.xml")
    except KeyError:
        return []
    root = ET.fromstring(xml)
    strings: list[str] = []
    for item in root.findall(f".//{{{_XLSX_MAIN_NS}}}si"):
        strings.append(_xlsx_join_text(item))
    return strings


def _xlsx_sheet_refs(archive: ZipFile) -> list[tuple[str, str]]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = _xlsx_workbook_relationships(archive)
    sheets: list[tuple[str, str]] = []
    for sheet_index, sheet in enumerate(workbook.findall(f".//{{{_XLSX_MAIN_NS}}}sheet"), 1):
        name = sheet.attrib.get("name") or f"Sheet{sheet_index}"
        relationship_id = sheet.attrib.get(f"{{{_XLSX_OFFICE_REL_NS}}}id")
        target = relationships.get(relationship_id or "")
        sheet_path = _xlsx_resolve_part("xl/workbook.xml", target) if target else f"xl/worksheets/sheet{sheet_index}.xml"
        sheets.append((name, sheet_path))
    return sheets


def _xlsx_workbook_relationships(archive: ZipFile) -> dict[str, str]:
    try:
        xml = archive.read("xl/_rels/workbook.xml.rels")
    except KeyError:
        return {}
    root = ET.fromstring(xml)
    relationships: dict[str, str] = {}
    for relationship in root.findall(f".//{{{_PACKAGE_REL_NS}}}Relationship"):
        relationship_id = relationship.attrib.get("Id")
        target = relationship.attrib.get("Target")
        if relationship_id and target:
            relationships[relationship_id] = target
    return relationships


def _xlsx_resolve_part(base_part: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def _xlsx_rows(xml: bytes, shared_strings: list[str]) -> list[list[str]]:
    root = ET.fromstring(xml)
    rows: list[list[str]] = []
    for row in root.findall(f".//{{{_XLSX_MAIN_NS}}}sheetData/{{{_XLSX_MAIN_NS}}}row"):
        values: list[str] = []
        for cell in row.findall(f"{{{_XLSX_MAIN_NS}}}c"):
            column_index = _xlsx_column_index(cell.attrib.get("r", ""))
            while len(values) < column_index:
                values.append("")
            values.append(_xlsx_cell_value(cell, shared_strings))
        while values and not values[-1]:
            values.pop()
        if any(values):
            rows.append(values)
    return rows


def _xlsx_column_index(reference: str) -> int:
    letters = "".join(char for char in reference if char.isalpha())
    if not letters:
        return 0
    index = 0
    for char in letters.upper():
        index = index * 26 + (ord(char) - ord("A") + 1)
    return max(0, index - 1)


def _xlsx_cell_value(cell: ET.Element, shared_strings: list[str]) -> str:
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        inline = cell.find(f"{{{_XLSX_MAIN_NS}}}is")
        return _xlsx_clean_value(_xlsx_join_text(inline) if inline is not None else "")

    value = cell.findtext(f"{{{_XLSX_MAIN_NS}}}v") or ""
    if cell_type == "s":
        try:
            return _xlsx_clean_value(shared_strings[int(value)])
        except (ValueError, IndexError):
            return ""
    if cell_type == "b":
        return "TRUE" if value == "1" else "FALSE"
    if cell_type == "str":
        return _xlsx_clean_value(value)
    return _xlsx_clean_value(value)


def _xlsx_join_text(element: ET.Element | None) -> str:
    if element is None:
        return ""
    return "".join(element.itertext())


def _xlsx_clean_value(value: str) -> str:
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"\s+", " ", line).strip() for line in value.split("\n")]
    return "\n".join(line for line in lines if line)


def _xlsx_rows_to_markdown(rows: list[list[str]]) -> str:
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    if not any(header):
        header = [f"列{index}" for index in range(1, width + 1)]
    body_rows = normalized[1:]
    lines = [
        "| " + " | ".join(_markdown_table_cell(value or f"列{index}") for index, value in enumerate(header, 1)) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    for row in body_rows:
        lines.append("| " + " | ".join(_markdown_table_cell(value) for value in row) + " |")
    return "\n".join(lines)


def _markdown_table_cell(value: str) -> str:
    return value.replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


def _parse_html(file_path: str) -> str:
    """Extract text from HTML using markdownify (HTML → Markdown)."""
    content = Path(file_path).read_text(encoding="utf-8", errors="replace")
    try:
        from readability import Document  # type: ignore

        content = Document(content).summary()
    except ImportError:
        pass
    try:
        import markdownify  # type: ignore
        return markdownify.markdownify(content, heading_style="ATX")
    except ImportError:
        # Fallback: strip tags with a simple regex
        import re
        text = re.sub(r"<[^>]+>", " ", content)
        return re.sub(r"\s+", " ", text).strip()


def _parse_text(file_path: str) -> str:
    """Read a plain text / markdown file."""
    return Path(file_path).read_text(encoding="utf-8", errors="replace")


def _parse_image(file_path: str, use_ocr: bool = False) -> str:
    """Extract text from an image when OCR is explicitly enabled."""
    if not use_ocr:
        raise ValueError("Image parsing requires ParseConfig(use_ocr=True).")
    try:
        from rapidocr_onnxruntime import RapidOCR  # type: ignore
    except ImportError as exc:
        raise ImportError(
            "rapidocr-onnxruntime is required for image OCR. Install nexagent-core[advanced-parsing]."
        ) from exc

    engine = RapidOCR()
    result, _ = engine(file_path)
    if not result:
        return ""
    return "\n".join(str(item[1]) for item in result if len(item) > 1 and item[1])


def _base_metadata(path: Path, content: str, parser: str) -> dict:
    stat = path.stat()
    return {
        "source": str(path),
        "filename": path.name,
        "suffix": path.suffix.lower(),
        "parser": parser,
        "bytes": stat.st_size,
        "content_chars": len(content),
        "content_sha256": hashlib.sha256(content.encode("utf-8", errors="ignore")).hexdigest(),
    }


_XLSX_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_XLSX_OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
